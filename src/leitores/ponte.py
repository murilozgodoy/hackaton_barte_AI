"""Leitor do e-mail da Ponte Câmbio (.eml com relatório .xlsx em base64).

O anexo é lido em memória. Linhas de total no fim do relatório ("Total de
operações:", "Spread total R$:") não viram operação: ficam em
attrs["totais_declarados"] para conferência.
"""
from __future__ import annotations

import email
import io
from datetime import date, datetime
from email import policy
from pathlib import Path

import openpyxl
import pandas as pd

from .comum import ErroLeitura, data_br, decimal_de_celula, normalizar, normalizar_cnpj

_MAPA = {
    "NO OPERACAO": "num_operacao",
    "DATA CONTRATACAO": "data_contratacao",
    "DATA LIQUIDACAO": "data_liquidacao",
    "CNPJ CLIENTE": "cnpj_original",
    "CLIENTE": "cliente",
    "MOEDA": "moeda",
    "VALOR ME": "valor_me",
    "TAXA CLIENTE": "taxa_cliente",
    "PTAX REF.": "ptax",
    "VALOR R$": "valor_brl",
    "SPREAD R$": "spread_brl",
    "SPREAD %": "spread_frac",
    "SPREAD ESTIMADO R$ (D-1)": "spread_estimado_brl",
    "IOF R$": "iof_brl",
    "STATUS COMPLIANCE": "status_compliance",
    "OBSERVACAO": "observacao",
}
_NUMERICAS = ["valor_me", "taxa_cliente", "ptax", "valor_brl", "spread_brl", "spread_frac",
              "spread_estimado_brl", "iof_brl"]


def _anexo_xlsx(caminho: Path) -> tuple[bytes, str, dict]:
    msg = email.message_from_bytes(caminho.read_bytes(), policy=policy.default)
    meta = {"assunto": str(msg.get("Subject", "")), "data_email": str(msg.get("Date", "")),
            "remetente": str(msg.get("From", ""))}
    for parte in msg.walk():
        nome = parte.get_filename() or ""
        if nome.lower().endswith(".xlsx") or parte.get_content_type().endswith("spreadsheetml.sheet"):
            return parte.get_payload(decode=True), nome, meta
    raise ErroLeitura(f"{caminho.name}: nenhum anexo .xlsx encontrado no e-mail")


def _eh_data(v) -> bool:
    if isinstance(v, (datetime, date)):
        return True
    try:
        return data_br(v) is not None
    except ErroLeitura:
        return False


def ler_ponte(caminho: Path, aba: str) -> pd.DataFrame:
    caminho = Path(caminho)
    conteudo, nome_anexo, meta = _anexo_xlsx(caminho)
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
    ws = wb[aba] if aba in wb.sheetnames else next(
        (wb[n] for n in wb.sheetnames if normalizar(n) == normalizar(aba)), None)
    if ws is None:
        raise ErroLeitura(f"{nome_anexo}: aba {aba!r} não encontrada. Abas: {wb.sheetnames}")
    linhas = list(ws.iter_rows(values_only=True))

    idx_cab = next((i for i, row in enumerate(linhas)
                    if "NO OPERACAO" in {normalizar(v) for v in row}), None)
    if idx_cab is None:
        raise ErroLeitura(f"{nome_anexo}/{ws.title}: cabeçalho com 'Nº Operação' não encontrado")
    colunas = {}
    for j, v in enumerate(linhas[idx_cab]):
        chave = _MAPA.get(normalizar(v))
        if chave:
            colunas[chave] = j
    faltando = [c for c in _MAPA.values() if c not in colunas]
    if faltando:
        raise ErroLeitura(f"{nome_anexo}/{ws.title}: colunas ausentes: {faltando}")

    registros, totais, ignoradas = [], {}, []
    for i in range(idx_cab + 1, len(linhas)):
        row = linhas[i]
        n_excel = i + 1
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        op = row[colunas["num_operacao"]]
        contratacao = row[colunas["data_contratacao"]]
        if op and _eh_data(contratacao):
            r = {k: row[j] for k, j in colunas.items()}
            registros.append(_converter(r, n_excel, caminho.name))
        elif isinstance(op, str) and op.strip().endswith(":"):
            valores = [v for v in row[colunas["num_operacao"] + 1:] if v is not None]
            totais[normalizar(op)] = decimal_de_celula(valores[0]) if valores else None
        else:
            ignoradas.append({"linha": n_excel, "conteudo": [v for v in row if v is not None]})

    df = pd.DataFrame(registros)
    df["data_contratacao"] = pd.to_datetime(df["data_contratacao"])
    df["data_liquidacao"] = pd.to_datetime(df["data_liquidacao"])
    df.attrs.update(arquivo=caminho.name, anexo=nome_anexo, aba=ws.title,
                    totais_declarados=totais, linhas_ignoradas=ignoradas, **meta)
    return df


def _converter(r: dict, linha: int, arquivo: str) -> dict:
    out = {
        "num_operacao": str(r["num_operacao"]).strip(),
        "data_contratacao": data_br(r["data_contratacao"]),
        "data_liquidacao": data_br(r["data_liquidacao"]),
        "cnpj": normalizar_cnpj(r["cnpj_original"]),
        "cnpj_original": r["cnpj_original"],
        "cliente": r["cliente"],
        "moeda": r["moeda"],
        "status_compliance": normalizar(r["status_compliance"]) or None,
        "observacao": (str(r["observacao"]).strip() if r["observacao"] else None),
    }
    for k in _NUMERICAS:
        out[k] = decimal_de_celula(r[k])
    # Spread % vem como fração (0.01501); em % fica comparável à planilha de taxas.
    out["spread_pct"] = out["spread_frac"] * 100 if out["spread_frac"] is not None else None
    out.update(linha_planilha=linha, ref_origem=out["num_operacao"], arquivo_origem=arquivo)
    return out
