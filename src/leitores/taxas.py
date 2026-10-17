"""Leitor da planilha de taxas contratadas (mantida à mão pelo financeiro).

Devolve formato longo: uma linha por (linha da planilha x produto), com a célula
de origem. Produtos são deduzidos do cabeçalho:
  "Débito" -> DEBITO; "Crédito à vista" -> CREDITO A VISTA; "Créd. 2-6x" ->
  CREDITO PARCELADO com parcelas 2 a 6; "Pix" -> PIX; "Antecipação" -> ANTECIPACAO;
  "Spread remessa" -> SPREAD REMESSA.
Os nomes normalizados batem com a modalidade/tipo da Orbe, sem tabela de tradução.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from .comum import (
    ErroLeitura, normalizar, normalizar_cnpj, normalizar_ec, parse_vigencia, taxa_pct_de_celula,
)

_RE_PARCELADO = re.compile(r"^CRED(ITO)?\.? *(PARCELADO)? *(\d+) *- *(\d+) *X$")
_PRODUTOS_SIMPLES = {"DEBITO", "CREDITO A VISTA", "PIX", "ANTECIPACAO", "SPREAD REMESSA"}
_COLUNAS_CADASTRO = {"LOJISTA": "lojista_planilha", "CNPJ": "cnpj", "EC ORBE": "ec_orbe",
                     "VIGENCIA": "vigencia_texto", "OBS": "obs"}


def _produto(cabecalho: str):
    if cabecalho in _PRODUTOS_SIMPLES:
        return cabecalho, None, None
    m = _RE_PARCELADO.match(cabecalho)
    if m:
        return "CREDITO PARCELADO", int(m.group(3)), int(m.group(4))
    return None


def _achar_aba(wb, aba: str):
    if aba in wb.sheetnames:
        return wb[aba]
    for nome in wb.sheetnames:
        if normalizar(nome) == normalizar(aba):
            return wb[nome]
    raise ErroLeitura(f"Aba {aba!r} não encontrada. Abas no arquivo: {wb.sheetnames}")


def ler_taxas(caminho: Path, aba: str) -> pd.DataFrame:
    caminho = Path(caminho)
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = _achar_aba(wb, aba)
    linhas = list(ws.iter_rows(values_only=True))

    # cabeçalho = primeira linha que tem "Lojista" e "CNPJ" (a linha acima traz só grupos)
    idx_cab = next((i for i, row in enumerate(linhas)
                    if {"LOJISTA", "CNPJ"} <= {normalizar(v) for v in row}), None)
    if idx_cab is None:
        raise ErroLeitura(f"{caminho.name}/{ws.title}: cabeçalho com 'Lojista' e 'CNPJ' não encontrado")
    cab = [normalizar(v) for v in linhas[idx_cab]]

    col_cadastro, col_produto, ignoradas = {}, {}, []
    for j, nome in enumerate(cab):
        if not nome:
            continue
        if nome in _COLUNAS_CADASTRO:
            col_cadastro[_COLUNAS_CADASTRO[nome]] = j
        elif _produto(nome):
            col_produto[j] = _produto(nome)
        else:
            ignoradas.append(nome)
    faltando = [c for c in ("lojista_planilha", "cnpj", "vigencia_texto") if c not in col_cadastro]
    if faltando:
        raise ErroLeitura(f"{caminho.name}/{ws.title}: colunas ausentes: {faltando}")

    registros, linhas_ignoradas = [], []
    for i in range(idx_cab + 1, len(linhas)):
        row = linhas[i]
        n_excel = i + 1
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        get = lambda campo: row[col_cadastro[campo]] if campo in col_cadastro else None
        cnpj = normalizar_cnpj(get("cnpj"))
        if cnpj is None:
            # ex.: nota de rodapé "* spread em branco = padrão"
            texto = " | ".join(str(v) for v in row if v is not None)
            linhas_ignoradas.append({"linha": n_excel, "conteudo": texto, "motivo": "sem CNPJ"})
            continue
        vig_de, vig_ate = parse_vigencia(get("vigencia_texto"))
        try:
            ec = normalizar_ec(get("ec_orbe"))
        except ErroLeitura as e:
            raise ErroLeitura(f"{caminho.name}/{ws.title} linha {n_excel}: {e}") from None
        base = {
            "linha_planilha": n_excel,
            "lojista_planilha": (str(get("lojista_planilha")).strip() if get("lojista_planilha") else None),
            "cnpj": cnpj,
            "ec_orbe": ec,
            "vigencia_texto": get("vigencia_texto"),
            "vigencia_de": vig_de,
            "vigencia_ate": vig_ate,
            "obs": get("obs"),
            "arquivo_origem": caminho.name,
        }
        for j, (produto, p_de, p_ate) in col_produto.items():
            valor = row[j]
            try:
                pct, situacao = taxa_pct_de_celula(valor)
            except ErroLeitura as e:
                raise ErroLeitura(f"{caminho.name}/{ws.title} {get_column_letter(j + 1)}{n_excel}: {e}") from None
            registros.append({
                **base,
                "produto": produto,
                "parcelas_de": p_de,
                "parcelas_ate": p_ate,
                "taxa_pct": pct,
                "taxa_situacao": situacao,
                "taxa_original": valor,
                "celula": f"{get_column_letter(j + 1)}{n_excel}",
                "ref_origem": f"{ws.title}!{get_column_letter(j + 1)}{n_excel}",
            })

    df = pd.DataFrame(registros)
    df["vigencia_de"] = pd.to_datetime(df["vigencia_de"])
    df["vigencia_ate"] = pd.to_datetime(df["vigencia_ate"])
    df["parcelas_de"] = df["parcelas_de"].astype("Int64")
    df["parcelas_ate"] = df["parcelas_ate"].astype("Int64")
    df.attrs.update(arquivo=caminho.name, aba=ws.title, linha_cabecalho=idx_cab + 1,
                    linhas_ignoradas=linhas_ignoradas, colunas_ignoradas=ignoradas)
    return df
