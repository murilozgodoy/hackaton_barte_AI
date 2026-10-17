"""Outubro sintético: um pacote completo do mês seguinte, gerado a partir dos arquivos de setembro.

Serve para o teste de férias (evals/teste_ferias.md) e para a robustez do motor em um mês que ainda não existe:
nenhum número de setembro muda, só o que o roteiro pede. O que o pacote traz, além das datas deslocadas +1 mês:

  Orbe      EC novo fora da planilha (12 transações de um lojista passam a ser do EC 9999999999 "LOJA NOVA LTDA");
            Bridge já na taxa nova (crédito à vista a 2,19%); Casa do Estudante zerada (a Orbe aplicou a isenção:
            taxa 0 em todas as transações); uma venda parcelada em 13x (acima das faixas da planilha);
            a Atlas continua transacionando (a carteira de setembro entra via --anterior).
  Taxas     a mesma planilha de setembro (a Bridge já tem a vigência nova; o EC novo não está nela).
  Ponte     operações com Nº PC-1026-*, competência 10/2026, mais 1 linha de ajuste no formato que a Helena descreveu
            (Nº próprio, Spread R$ negativo, observação "AJUSTE REF PC-0926-0016": excesso da Mundo Aberto devolvido,
            pedido a da contestação de setembro) e 2 operações de um CNPJ novo.
  Plano     o mesmo.
  Aurora    datas +1 mês (folha e impostos "do mês anterior" passam a 09/2026); saldos iguais.
  Meridiano datas e FITIDs +1 mês.

Uso: python -m tests.outubro_sintetico <pasta de destino> [<pasta ou zip de setembro>]
"""
from __future__ import annotations

import csv
import email
import io
import re
import shutil
import sys
import zipfile
from datetime import date
from decimal import Decimal, ROUND_DOWN
from email import policy
from email.message import EmailMessage
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
D = Decimal
EC_NOVO, NOME_EC_NOVO = "9999999999", "LOJA NOVA LTDA"
CNPJ_NOVO_PONTE, NOME_CNPJ_NOVO = "99999999999999", "LOJA NOVA CAMBIO LTDA"
EC_BRIDGE, EC_CASA, EC_ATLAS = "1075624862", "1018325654", "1076801732"
EC_TROCADO = "1043300981"          # Novo Horizonte: 12 transações viram do EC novo
TAXA_BRIDGE_NOVA = D("2.19")


def mes_seguinte(d: date) -> date:
    ano, mes = (d.year, d.month + 1) if d.month < 12 else (d.year + 1, 1)
    ultimo = [31, 29 if ano % 4 == 0 and (ano % 100 != 0 or ano % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][mes - 1]
    return date(ano, mes, min(d.day, ultimo))


def _br(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _shift_br(texto: str) -> str:
    """dd/mm/aaaa -> mês seguinte; mm/aaaa (ex.: 'ISS 08/2026') -> mês seguinte."""
    def f(m):
        return _br(mes_seguinte(date(int(m.group(3)), int(m.group(2)), int(m.group(1)))))
    texto = re.sub(r"\b(\d{2})/(\d{2})/(\d{4})\b", f, texto)
    def g(m):
        d = mes_seguinte(date(int(m.group(2)), int(m.group(1)), 1))
        return f"{d.month:02d}/{d.year}"
    return re.sub(r"(?<!/)\b(\d{2})/(\d{4})\b", g, texto)


def _pasta_setembro(origem: Path | None) -> Path:
    if origem is None:
        candidatos = sorted(RAIZ.glob("Tarvo_Data_Pack_*.zip"))
        if not candidatos:
            raise SystemExit("não achei o zip de setembro na raiz; passe a pasta ou o zip como 2º argumento")
        origem = candidatos[0]
    origem = Path(origem)
    if origem.suffix.lower() == ".zip":
        # extrai dentro de data/<nome do zip>/ (com ou sem pasta interna no zip); os arquivos são achados por rglob
        destino = Path(RAIZ / "data" / origem.stem)
        if not any(destino.rglob("01_orbe*")):
            destino.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(origem) as z:
                z.extractall(destino)
        return destino
    return origem


def _um(pasta: Path, padrao: str) -> Path:
    achados = sorted(pasta.rglob(padrao))
    if len(achados) != 1:
        raise SystemExit(f"esperava um arquivo {padrao} em {pasta}, achei {len(achados)}")
    return achados[0]


# --------------------------------------------------------------------------- Orbe

def _orbe(origem: Path, destino: Path) -> None:
    texto = origem.read_text(encoding="utf-8-sig")
    linhas = list(csv.reader(io.StringIO(texto), delimiter=";"))
    cab, corpo = linhas[0], linhas[1:]
    idx = {c: i for i, c in enumerate(cab)}
    trocadas = 0
    out = []
    for campos in corpo:
        if not any(campos):
            continue
        c = list(campos)
        c[idx["id_transacao"]] = "TY" + c[idx["id_transacao"]][2:]
        c[idx["id_referencia"]] = ("TY" + c[idx["id_referencia"]][2:]) if c[idx["id_referencia"]] else ""
        c[idx["data_movimento"]] = _shift_br(c[idx["data_movimento"]])
        c[idx["data_venda_original"]] = _shift_br(c[idx["data_venda_original"]]) if c[idx["data_venda_original"]] else ""
        ec = c[idx["cod_estabelecimento"]]
        bruto = D(c[idx["valor_bruto"]].replace(".", "").replace(",", "."))
        taxa = D(c[idx["taxa_aplicada_pct"]].replace(",", "."))
        if ec == EC_TROCADO and trocadas < 12:
            c[idx["cod_estabelecimento"]], c[idx["nome_estabelecimento"]] = EC_NOVO, NOME_EC_NOVO
            trocadas += 1
        if ec == EC_BRIDGE and c[idx["modalidade"]].upper().startswith("CREDITO A VISTA") and taxa == D("2.59"):
            taxa = TAXA_BRIDGE_NOVA
        if ec == EC_CASA:
            taxa = D("0")
        valor_taxa = (bruto * taxa / 100).quantize(D("0.01"), rounding=ROUND_DOWN) if bruto >= 0 else \
            -((-bruto) * taxa / 100).quantize(D("0.01"), rounding=ROUND_DOWN)
        c[idx["taxa_aplicada_pct"]] = f"{taxa:.2f}".replace(".", ",")
        c[idx["valor_taxa"]] = _num(valor_taxa)
        c[idx["valor_liquido"]] = _num(bruto - valor_taxa)
        out.append(c)
    # uma venda em 13x da Rota Global (EC da primeira linha com CREDITO PARCELADO que não seja Casa/Atlas)
    modelo = next(c for c in out if c[idx["modalidade"]].upper() == "CREDITO PARCELADO"
                  and c[idx["cod_estabelecimento"]] not in (EC_CASA, EC_ATLAS, EC_NOVO))
    extra = list(modelo)
    extra[idx["id_transacao"]], extra[idx["id_referencia"]] = "TY99999913", ""
    extra[idx["qtd_parcelas"]], extra[idx["tipo_lancamento"]] = "13", "VENDA"
    extra[idx["data_venda_original"]] = ""
    bruto = D("1300.00")
    taxa = D(extra[idx["taxa_aplicada_pct"]].replace(",", "."))
    vt = (bruto * taxa / 100).quantize(D("0.01"), rounding=ROUND_DOWN)
    extra[idx["valor_bruto"]], extra[idx["valor_taxa"]], extra[idx["valor_liquido"]] = _num(bruto), _num(vt), _num(bruto - vt)
    out.append(extra)
    with open(destino, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\n")
        w.writerow(cab)
        w.writerows(out)


def _num(v: Decimal) -> str:
    s = f"{v.quantize(D('0.01')):,.2f}"          # 1,234.56
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


# --------------------------------------------------------------------------- Ponte (e-mail com xlsx)

def _ponte(origem: Path, destino: Path) -> None:
    msg = email.message_from_bytes(origem.read_bytes(), policy=policy.default)
    corpo = ""
    xlsx = nome = None
    for parte in msg.walk():
        if parte.get_content_type() == "text/plain":
            corpo += parte.get_content()
        if (parte.get_filename() or "").lower().endswith(".xlsx"):
            xlsx, nome = parte.get_payload(decode=True), parte.get_filename()
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    ws = wb["Operações"]
    linhas = list(ws.iter_rows())
    idx_cab = next(i for i, row in enumerate(linhas) if any(str(c.value or "").strip().startswith("Nº Opera") for c in row))
    cab = {str(c.value).strip(): j for j, c in enumerate(linhas[idx_cab]) if c.value}
    col = lambda nome_col: cab[next(k for k in cab if k.startswith(nome_col))]
    for row in linhas[:idx_cab]:
        for c in row:
            if isinstance(c.value, str):
                c.value = _shift_br(c.value)
    n_ops, spread_total, trocadas, ultima = 0, D("0"), 0, idx_cab
    for i in range(idx_cab + 1, len(linhas)):
        row = linhas[i]
        op = row[col("Nº Opera")].value
        if isinstance(op, str) and op.startswith("PC-"):
            row[col("Nº Opera")].value = op.replace("PC-0926-", "PC-1026-")
            for nome_col in ("Data Contrata", "Data Liquida"):
                c = row[col(nome_col)]
                if c.value is not None:
                    d = c.value.date() if hasattr(c.value, "date") else c.value
                    c.value = mes_seguinte(d)
            if trocadas < 2 and row[col("Cliente")].value and str(row[col("Cliente")].value).startswith("MAPLE"):
                row[col("CNPJ Cliente")].value, row[col("Cliente")].value = CNPJ_NOVO_PONTE, NOME_CNPJ_NOVO
                trocadas += 1
            n_ops += 1
            spread_total += D(str(row[col("Spread R$")].value))
            ultima = i
        elif isinstance(op, str) and op.strip().endswith(":"):
            if "opera" in op.lower():
                celula_total_ops = row[col("Nº Opera") + 1]
            else:
                celula_total_spread = row[col("Nº Opera") + 1]
    # ajuste da contestação de setembro no formato da Ponte (Helena à Desbravadores 17/10 11h33): linha nova, Nº próprio,
    # Spread R$ negativo, observação com a operação original
    modelo = [c.value for c in linhas[ultima]]
    ajustes = [
        ("PC-1026-0901", date(2026, 10, 5), "17217888000111", "MUNDO ABERTO INTERCÂ", D("-45.00"), D("-0.003"),
         "AJUSTE REF PC-0926-0016"),
    ]
    for k, (num, d, cnpj, cliente, spread, frac, obs) in enumerate(ajustes):
        ws.insert_rows(ultima + 2 + k)
        valores = list(modelo)
        valores[col("Nº Opera")], valores[col("Data Contrata")], valores[col("Data Liquida")] = num, d, mes_seguinte(d)
        valores[col("CNPJ Cliente")], valores[col("Cliente")], valores[col("Moeda")] = cnpj, cliente, "USD"
        valores[col("Valor ME")], valores[col("Taxa Cliente")], valores[col("PTAX")] = 0, 0, 0
        valores[col("Valor R$")], valores[col("Spread R$")], valores[col("Spread %")] = 0, float(spread), float(frac)
        valores[col("Spread Estimado")], valores[col("IOF")], valores[col("Status")], valores[col("Observa")] = 0, 0, "APROVADO", obs
        for j, v in enumerate(valores):
            ws.cell(row=ultima + 2 + k, column=j + 1, value=v)
        n_ops += 1
        spread_total += spread
    # totais declarados (as células foram deslocadas pelo insert_rows: localizar de novo)
    for row in ws.iter_rows():
        v = row[col("Nº Opera")].value
        if isinstance(v, str) and v.strip().endswith(":"):
            row[col("Nº Opera") + 1].value = n_ops if "opera" in v.lower() else float(spread_total)
    buf = io.BytesIO()
    wb.save(buf)

    novo = EmailMessage()
    novo["From"] = str(msg["From"])
    novo["To"] = str(msg["To"])
    novo["Subject"] = _shift_br(str(msg["Subject"]))
    novo["Date"] = "Mon, 02 Nov 2026 07:15:03 -0300"
    novo.set_content(_shift_br(corpo))
    novo.add_attachment(buf.getvalue(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        filename=nome.replace("092026", "102026"))
    destino.write_bytes(novo.as_bytes())


# --------------------------------------------------------------------------- bancos

def _aurora(origem: Path, destino: Path) -> None:
    raw = origem.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            texto = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    destino.write_text(_shift_br(texto), encoding=enc)


def _meridiano(origem: Path, destino: Path) -> None:
    raw = origem.read_bytes()
    for enc in ("utf-8", "cp1252"):
        try:
            texto = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    def f(m):
        d = mes_seguinte(date(int(m.group(1)), int(m.group(2)), int(m.group(3))))
        return d.strftime("%Y%m%d")
    texto = re.sub(r"(?<=>)(\d{4})(\d{2})(\d{2})(?=\d{6}\[)", f, texto)     # DTPOSTED, DTSTART, DTEND, DTASOF, DTSERVER
    def g(m):   # FITID = MRD + AAAAMM + sequência
        d = mes_seguinte(date(int(m.group(1)), int(m.group(2)), 1))
        return f"MRD{d.year}{d.month:02d}"
    texto = re.sub(r"MRD(\d{4})(\d{2})(?=\d{4}\b)", g, texto)
    texto = _shift_br(texto)   # "PARC 09/2026" etc.
    destino.write_text(texto, encoding=enc)


# --------------------------------------------------------------------------- gerar

def gerar(destino: Path, setembro: Path | None = None) -> Path:
    """Gera o pacote de outubro em `destino` (pasta); devolve a pasta."""
    pasta = _pasta_setembro(setembro)
    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    _orbe(_um(pasta, "01_orbe_*.csv"), destino / "01_orbe_relatorio_comissoes_out2026.csv")
    shutil.copy(_um(pasta, "02_tabela_taxas*.xlsx"), destino / "02_tabela_taxas_contratadas.xlsx")
    _ponte(_um(pasta, "03_email_ponte*.eml"), destino / "03_email_ponte_cambio_relatorio_out2026.eml")
    shutil.copy(_um(pasta, "04_plano_de_contas*.csv"), destino / "04_plano_de_contas_fenix_erp.csv")
    _aurora(_um(pasta, "05a_extrato_banco_aurora*.csv"), destino / "05a_extrato_banco_aurora_cc4471_out2026.csv")
    _meridiano(_um(pasta, "05b_extrato_banco_meridiano*.ofx"), destino / "05b_extrato_banco_meridiano_cc90233_out2026.ofx")
    return destino


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    print(gerar(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) > 2 else None))
